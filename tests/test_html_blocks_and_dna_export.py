"""Draft DNA on the fields drafting screen — the seam that was missing.

Until this work the ONLY producer of Draft DNA's role-tagged blocks was the LLM
authoring path. Every other way an advocate actually gets a draft — the fields
screen at /draft/template/<id> (all 50 reviewed types), the canonical templates,
and the deterministic floors in from_prompt — produced HTML, so "Download .docx
in my own format" was unreachable from the screens most advocates use, and
GET /api/draft/<id>/docx answered 409 whenever the model had been slow or capped.

What these tests hold down:

  1. The HTML→blocks converter recovers the FULL role vocabulary for EVERY
     reviewed type, in every language the engine renders — because the mapping is
     structural (the engine's own CSS classes), never Hindi words. This is the
     universality property: a Chennai advocate drafting in English and a Pune
     advocate drafting in Marathi go down the identical path.
  2. The .docx genuinely comes out DIFFERENT for two different advocates — same
     document, his page size, his margins, his typeface. If this test can't tell
     two advocates apart, Draft DNA is decoration.
  3. With no DNA the action still works and says so honestly (standard format).
  4. What he EDITED on the canvas is what he downloads.
  5. The AI change box only turns knobs the reviewed template actually has.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from fastapi.testclient import TestClient

from headnote.drafter import html_blocks as HB
from headnote.drafter import layout_template as LT
from headnote.drafter import template_adapter as TA

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# One representative of each shape the engine can emit: criminal application with
# a bundle + affidavit table, a civil plaint, a family petition, a writ, and the
# one type with no numbered grounds at all.
SAMPLE_TYPES = ["bail_sessions", "recovery_suit", "maintenance",
                "writ_petition", "vakalatnama", "cheque_138"]


# --------------------------------------------------------------- the converter

# A sworn affidavit and a demand notice are not court filings: there is no court
# and no opposite party to name. They still need every other block.
_NO_CAUSE_TITLE = {"general_affidavit", "legal_notice"}


@pytest.mark.parametrize("tid", sorted(TA.CANONICAL_MAP))
def test_every_reviewed_type_yields_layout_blocks(tid):
    """All 50 types, not a curated few. A type that yields no blocks is a type
    whose advocate silently cannot get a .docx in his own format."""
    blocks = HB.from_document_html(TA.document(tid, {}, "hi"))
    assert not HB.is_empty(blocks), f"{tid} produced no layout blocks"
    roles = {r for r, _ in blocks}
    if tid not in _NO_CAUSE_TITLE:
        # The cause title is what makes a filing filable — if these are missing
        # the .docx is an undifferentiated wall of body text.
        for essential in ("court", "caseno", "applicant", "respondent"):
            assert essential in roles, f"{tid} lost the {essential!r} block"
    assert roles <= set(LT.ROLES) | {"table", "page_break"}, \
        f"{tid} emitted a role Draft DNA cannot format: {roles - set(LT.ROLES)}"


@pytest.mark.parametrize("tid", sorted(TA.CANONICAL_MAP))
def test_no_type_starts_with_a_page_break(tid):
    """A leading break is a blank first page in the filed document. Some builders
    nest one sheet wrapper inside another, which is how this got in."""
    blocks = HB.from_document_html(TA.document(tid, {}, "hi"))
    assert blocks[0] != HB.PAGE_BREAK, f"{tid} would open on a blank page"


@pytest.mark.parametrize("tid", SAMPLE_TYPES)
@pytest.mark.parametrize("lang", ["hi", "en"])
def test_converter_is_language_neutral(tid, lang):
    """The SAME structure comes back whatever language the document is in.

    This is the property that makes Draft DNA work for advocates outside the one
    practice it was built against: the mapping reads the engine's CSS classes, so
    it never has to know a single word of the language being drafted in.
    """
    blocks = HB.from_document_html(TA.document(tid, {}, lang))
    assert not HB.is_empty(blocks)
    roles_here = [r for r, _ in blocks]
    roles_hi = [r for r, _ in HB.from_document_html(TA.document(tid, {}, "hi"))]

    # The role VOCABULARY and the order of the structural spine must match. The
    # per-role COUNTS deliberately do not: some reviewed templates carry a
    # different number of grounds in English than in Hindi (writ_petition has 5
    # in Hindi and 3 in English, maintenance 8 vs 7). That asymmetry lives in the
    # legal content of the templates, not here — asserting equal counts would
    # wrongly blame this converter for it.
    assert set(roles_here) == set(roles_hi), \
        f"{tid} in {lang} lost a block type that Hindi has"
    dedupe = lambda rs: [r for i, r in enumerate(rs) if i == 0 or rs[i - 1] != r]
    assert dedupe(roles_here) == dedupe(roles_hi), \
        f"{tid} in {lang} produced a different document shape than in Hindi"


def test_grounds_keep_their_numbers_and_order():
    """A ground is argued in order and referred to by number in the verification
    ("paras 1 to 7"). Losing the numbering makes the affidavit wrong."""
    blocks = HB.from_document_html(TA.document("bail_sessions", {}, "hi"))
    grounds = [t for r, t in blocks if r == "ground"]
    assert len(grounds) >= 5
    assert [g.split(".")[0] for g in grounds[:5]] == ["1", "2", "3", "4", "5"]


def test_bundle_sheets_are_separated_by_a_page_break():
    """A filing is several sheets (Application · Affidavit). Without a break the
    affidavit starts halfway down the application's last page — not filable."""
    blocks = HB.from_document_html(TA.document("bail_sessions", {}, "hi"))
    assert ("page_break", "") in blocks


def test_tables_survive_as_tables():
    """The affidavit's particulars table is a table, not a paragraph of runs."""
    blocks = HB.from_document_html(TA.document("bail_sessions", {}, "hi"))
    tables = [c for r, c in blocks if r == "table"]
    assert tables and tables[0]["rows"], "the affidavit table was flattened"


def test_placeholders_are_kept_not_silently_dropped():
    """The greyed blanks on screen must be blanks in the .docx too. Dropping them
    hands the advocate a document that READS complete and is not."""
    html = TA.document("bail_sessions", {}, "hi")
    blocks = HB.from_document_html(html)
    text = " ".join(str(c) for r, c in blocks if r != "table")
    assert "…" in text or "___" in text or "...." in text or "पता" in text


def test_malformed_html_never_raises():
    for junk in ("", "   ", "<div>", "<p>hello", "not html at all"):
        assert isinstance(HB.from_document_html(junk), list)


# ------------------------------------------------------------ the .docx itself

def _docx_facts(data: bytes) -> dict:
    from docx import Document
    d = Document(io.BytesIO(data))
    s = d.sections[0]
    return {
        "page": (round(s.page_width.inches, 2), round(s.page_height.inches, 2)),
        "left_margin": round(s.left_margin.inches, 2),
        "fonts": {r.font.name for p in d.paragraphs for r in p.runs if r.font.name},
        "paras": len([p for p in d.paragraphs if p.text.strip()]),
        "tables": len(d.tables),
    }


def _capture(path: str) -> dict:
    with open(path, "rb") as fh:
        return LT.merge_templates([LT.capture_layout(fh.read())])


def test_two_advocates_get_visibly_different_documents():
    """The whole promise, as a measurement.

    Same filing, same facts — but one advocate files on 8.5x11 in Mangal and the
    other in Times New Roman. If Draft DNA is real these two files differ in page
    setup and typeface; if it is decoration they come out identical.
    """
    blocks = HB.from_document_html(TA.document("bail_sessions", {}, "hi"))
    a = _docx_facts(LT.render_into_layout(_capture("Headnote_Sample_bail_sessions.docx"), blocks))
    b = _docx_facts(LT.render_into_layout(_capture("Headnote_Bail_Hindi.docx"), blocks))
    none = _docx_facts(LT.render_into_layout(LT.standard_template(LT.default_font_for("hi")), blocks))

    assert a["fonts"] != b["fonts"], "two different advocates got the same typeface"
    assert none["fonts"] == {"Nirmala UI"}, "the no-DNA floor should be neutral"
    assert none["page"] == (8.27, 11.69), "the no-DNA floor should be plain A4"
    assert a["page"] != none["page"], "his captured page size was not applied"
    # and the document itself is intact in all three
    for facts in (a, b, none):
        assert facts["paras"] > 20 and facts["tables"] >= 1


def test_no_dna_still_produces_a_real_filable_docx():
    """"Download .docx" is never a dead action — worst case it is standard court
    format, and the screen says so."""
    blocks = HB.from_document_html(TA.document("maintenance", {}, "hi"))
    data = LT.render_into_layout(LT.standard_template("Nirmala UI"), blocks)
    assert zipfile.ZipFile(io.BytesIO(data)).namelist(), "not a valid .docx"
    assert _docx_facts(data)["paras"] > 10


# ------------------------------------------------------------------ the routes

@pytest.fixture()
def client(monkeypatch):
    from headnote.api.app import app
    from headnote.entitlements import CurrentUser
    from headnote.entitlements.auth import get_current_user, optional_user

    user = CurrentUser(id="00000000-0000-0000-0000-00000000d17a",
                       email="advocate@example.test", role="authenticated",
                       raw_claims={})
    saved = dict(app.dependency_overrides)
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[optional_user] = lambda: user
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(saved)


def _no_dna(monkeypatch):
    from headnote.drafter import style_profile as SP
    monkeypatch.setattr(SP, "load_layout", lambda uid: None)


def _with_dna(monkeypatch, path="Headnote_Sample_bail_sessions.docx"):
    from headnote.drafter import style_profile as SP
    tpl = _capture(path)
    monkeypatch.setattr(SP, "load_layout", lambda uid: tpl)
    return tpl


@pytest.mark.parametrize("tid", SAMPLE_TYPES)
def test_docx_route_works_for_every_type(client, monkeypatch, tid):
    _no_dna(monkeypatch)
    r = client.post("/api/draft/template-docx",
                    json={"doc_type": tid, "fields": {"applicant_name": "Ramu"}, "lang": "hi"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == DOCX_MIME
    assert r.headers["X-Headnote-Format"] == "standard"
    assert zipfile.ZipFile(io.BytesIO(r.content)).namelist()


def test_docx_route_uses_his_own_format_when_he_has_one(client, monkeypatch):
    _with_dna(monkeypatch)
    r = client.post("/api/draft/template-docx",
                    json={"doc_type": "bail_sessions", "fields": {}, "lang": "hi"})
    assert r.status_code == 200
    assert r.headers["X-Headnote-Format"] == "own"
    facts = _docx_facts(r.content)
    assert facts["fonts"] == {"Mangal"}
    assert facts["page"] == (8.5, 11.0)


def test_the_docx_carries_what_he_edited_on_the_canvas(client, monkeypatch):
    """He corrects a line on the page and downloads. He must get HIS correction,
    not a re-render from the form that silently discards it."""
    _no_dna(monkeypatch)
    edited = TA.document("bail_sessions", {}, "hi").replace(
        "बन्दी की ओर से", "CORRECTED BY THE ADVOCATE")
    r = client.post("/api/draft/template-docx",
                    json={"doc_type": "bail_sessions", "fields": {}, "lang": "hi",
                          "html": edited})
    assert r.status_code == 200
    from docx import Document
    body = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "CORRECTED BY THE ADVOCATE" in body


def test_export_refuses_rather_than_silently_downgrading_the_format(monkeypatch):
    """The bug that reads to an advocate as "Draft DNA doesn't work".

    These routes must NOT use `optional_user`: it returns None both when there is
    no token and when the token has merely EXPIRED — and a Supabase token expires
    after about an hour. Under `optional_user`, an advocate who has set up his
    format, leaves the tab open and comes back would download a .docx in standard
    format, with no error and nothing on screen to explain it. A 401 he can act on
    is the honest failure.
    """
    from headnote.api.app import app
    from headnote.drafter import api as drafter_api
    from headnote.entitlements.auth import get_current_user, optional_user
    from headnote.drafter import style_profile as SP

    PATHS = {"/api/draft/template-docx", "/api/draft/template-tweak"}

    # The wiring itself, which holds whatever order the suite runs in.
    checked = set()
    for route in drafter_api.router.routes:
        path = getattr(route, "path", "")
        full = "/api/draft" + path if not path.startswith("/api") else path
        if full not in PATHS:
            continue
        checked.add(full)
        deps = {d.call for d in route.dependant.dependencies}
        assert optional_user not in deps, f"{full} uses optional_user — see the docstring"
        assert get_current_user in deps, f"{full} must require a live session"
    assert checked == PATHS, f"route moved or renamed: only found {checked}"

    # And the behaviour, when auth is actually in production mode. Stated as a
    # precondition rather than asserted blindly: `headnote/entitlements/auth.py`
    # reads SUPABASE_URL at import time and elsewhere in this suite it is unset,
    # which switches auth into its local-dev synthetic-user mode for the rest of
    # the session. A green auth test that only holds in the right order is a lie.
    from headnote.entitlements import auth as _auth
    if not getattr(_auth, "SUPABASE_URL", ""):
        pytest.skip("auth is in local-dev mode (SUPABASE_URL unset) — 401 cannot be asserted")

    monkeypatch.setattr(SP, "load_layout", lambda uid: _capture("Headnote_Sample_bail_sessions.docx"))
    saved = dict(app.dependency_overrides)
    app.dependency_overrides.clear()          # nobody signed in — a stale token
    try:
        c = TestClient(app)
        for path, payload in (
            ("/api/draft/template-docx", {"doc_type": "bail_sessions", "fields": {}, "lang": "hi"}),
            ("/api/draft/template-tweak", {"doc_type": "bail_sessions", "fields": {},
                                           "prompt": "add a ground", "lang": "hi"}),
        ):
            r = c.post(path, json=payload)
            assert r.status_code == 401, \
                f"{path} answered {r.status_code} — it must refuse, not quietly " \
                f"hand him the wrong format"
    finally:
        app.dependency_overrides.update(saved)


def test_docx_route_refuses_an_unknown_type(client, monkeypatch):
    _no_dna(monkeypatch)
    r = client.post("/api/draft/template-docx",
                    json={"doc_type": "no_such_template", "fields": {}, "lang": "hi"})
    assert r.status_code == 400


def test_docx_route_refuses_an_empty_document(client, monkeypatch):
    _no_dna(monkeypatch)
    r = client.post("/api/draft/template-docx",
                    json={"doc_type": "bail_sessions", "fields": {}, "lang": "hi",
                          "html": "<div></div>"})
    assert r.status_code == 400


# ------------------------------------------------------- the AI change box

def test_tweak_moves_the_filing_to_another_forum(client):
    """The court is not a field — it is a different reviewed template. Asking for
    the High Court has to move the screen to the HC template, not edit a string."""
    r = client.post("/api/draft/template-tweak",
                    json={"doc_type": "bail_sessions", "fields": {"applicant_name": "Ramu"},
                          "prompt": "file it in the High Court instead",
                          "lang": "hi", "use_llm": False})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["doc_type"] == "bail_hc"
    assert d["fields"]["applicant_name"] == "Ramu", "his typed data was lost in the move"
    assert d["document"]


def test_tweak_never_invents_a_knob_the_template_does_not_have(client, monkeypatch):
    """The zero-fabrication rule, at the change box. A patch naming a field this
    reviewed template does not define must be REPORTED, not applied."""
    from headnote.drafter import prompt_tweak

    monkeypatch.setattr(prompt_tweak, "run_router", lambda spec, data, prompt: {
        "set": {"applicant_name": "Ramu", "invented_field": "nonsense"},
        "toggles": {"no_such_ground": True}, "variant": {}, "add_grounds": [], "note": "",
    })
    r = client.post("/api/draft/template-tweak",
                    json={"doc_type": "bail_sessions", "fields": {},
                          "prompt": "anything", "lang": "hi"})
    assert r.status_code == 200
    d = r.json()
    assert d["fields"].get("applicant_name") == "Ramu"      # the real knob turned
    assert "invented_field" not in d["fields"]              # the invented one did not
    assert "no_such_ground" not in d["fields"]
    ignored = [c for c in d["changelog"] if c.startswith("⚠")]
    assert len(ignored) == 2, "silently dropped a change instead of reporting it"


def test_tweak_refuses_an_unknown_type(client):
    r = client.post("/api/draft/template-tweak",
                    json={"doc_type": "nope", "fields": {}, "prompt": "x", "lang": "hi"})
    assert r.status_code == 404


def test_tweak_refuses_an_empty_prompt(client):
    r = client.post("/api/draft/template-tweak",
                    json={"doc_type": "bail_sessions", "fields": {}, "prompt": "  ",
                          "lang": "hi"})
    assert r.status_code == 400


# ----------------------------------------- the 409 that used to hit at the counter

def test_saved_draft_without_blocks_still_exports(client, monkeypatch):
    """A draft that landed on the canonical or skeleton floor has html but no
    `blocks` — those floors never produced any. That is exactly what happens when
    the model is slow, capped or down, and it used to make "Download .docx" answer
    409 on the day it was needed most. The roles are now recovered from the
    rendered document instead."""
    _no_dna(monkeypatch)
    from headnote.drafter import api as drafter_api

    class _Draft:
        id = "d1"
        user_id = "00000000-0000-0000-0000-00000000d17a"
        lang = "hi"
        title = "Bail"
        story_id = "bail"
        answers = {"html_hi": TA.document("bail_sessions", {}, "hi")}   # no "blocks"

    monkeypatch.setattr(drafter_api.storage, "get_draft", lambda did: _Draft())
    r = client.get("/api/draft/d1/docx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == DOCX_MIME
    assert _docx_facts(r.content)["paras"] > 20


# =========================================================================
# THE STYLE RAIL — page layout as a visible choice, and "match this document".
#
# Court practice is not uniform across India: an MP district court files on
# 8.5x14 legal paper, most High Court registries want double spacing and a
# binding margin, and an English commercial filing wants A4 in a serif. A single
# imposed house format serves none of them, so the layout is a choice on screen
# — and the .docx has to follow that choice exactly, or the advocate files
# something other than what he approved.
# =========================================================================

def _preset_facts(preset_id, lang="hi"):
    from headnote.drafter import layout_presets as LP
    blocks = HB.from_document_html(TA.document("bail_sessions", {}, lang))
    return _docx_facts(LT.render_into_layout(LP.build(preset_id, lang), blocks))


def test_every_preset_is_actually_a_different_document():
    """A rail of four options that all render the same file is decoration."""
    from headnote.drafter import layout_presets as LP

    seen = {}
    for pid in LP.PRESET_IDS:
        f = _preset_facts(pid)
        key = (f["page"], f["left_margin"])
        assert f["paras"] > 20, f"{pid} produced an empty document"
        seen[pid] = key
    assert seen["legal_district"][0] == (8.5, 14.0), "district legal size was not applied"
    assert seen["hc_double"][1] > seen["standard"][1], "the HC binding margin was not applied"
    assert len({v for v in seen.values()}) >= 3, "the presets are not meaningfully distinct"


def test_double_spacing_and_compact_really_change_the_spacing():
    """The two presets whose whole point is vertical rhythm."""
    from docx import Document
    from headnote.drafter import layout_presets as LP

    blocks = HB.from_document_html(TA.document("bail_sessions", {}, "hi"))

    def spacings(pid):
        d = Document(io.BytesIO(LT.render_into_layout(LP.build(pid, "hi"), blocks)))
        return {p.paragraph_format.line_spacing for p in d.paragraphs
                if p.paragraph_format.line_spacing}

    assert 2.0 in spacings("hc_double"), "double spacing was not applied"
    assert min(spacings("compact")) < 1.5, "compact did not tighten anything"


def test_presets_follow_his_script_not_a_hardcoded_font():
    """The same four presets have to serve an English filing in Chennai and a
    Marathi one in Pune — so the typeface comes from the language, never a
    Devanagari default baked into the preset."""
    assert _preset_facts("standard", "en")["fonts"] == {"Times New Roman"}
    assert _preset_facts("standard", "hi")["fonts"] == {"Nirmala UI"}


def test_reference_document_layout_is_captured_and_applied():
    """"Make it look like this one." The draft takes the reference's real page
    setup and typeface — not a guess, and not our house format."""
    from headnote.drafter import layout_presets as LP

    with open("Headnote_Bail_Hindi.docx", "rb") as fh:
        tpl = LP.capture_reference(fh.read())
    assert tpl and tpl.get("roles")
    blocks = HB.from_document_html(TA.document("bail_sessions", {}, "hi"))
    facts = _docx_facts(LT.render_into_layout(tpl, blocks))
    assert facts["fonts"] == {"Times New Roman"}
    assert facts["page"] == (8.5, 11.0)
    # roles the reference did not contain still render, in ITS typeface — a
    # half-formatted document would be worse than a fully standard one
    assert facts["paras"] > 20


def test_a_reference_never_overwrites_his_saved_draft_dna(monkeypatch):
    """A reference is a decision about ONE matter. Rewriting the format of every
    future draft from a single attachment is a much bigger change than he asked
    for, and he has an explicit place to set his own format."""
    from headnote.drafter import layout_presets as LP
    from headnote.drafter import style_profile as SP

    wrote = []
    monkeypatch.setattr(SP, "save_layout", lambda uid, layout: wrote.append(uid))
    with open("Headnote_Bail_Hindi.docx", "rb") as fh:
        LP.capture_reference(fh.read())
    assert wrote == [], "capturing a reference wrote to his saved Draft DNA"


@pytest.mark.parametrize("junk", [
    {"page": {"width": 9000, "height": -4}, "roles": {"ground": {"size": 99999}}},
    {"page": {"width": 0.5, "height": 0.5}, "roles": {"ground": {}}},
    {"page": {}, "roles": {}},
    {"roles": {"not_a_role": {"size": 12}}},
    "not a dict",
    None,
])
def test_a_layout_from_the_browser_is_never_trusted(junk):
    """The rail hands a captured layout back through the browser, so it is
    untrusted input. A 900-inch page or a negative margin must be dropped, not
    passed to python-docx to raise or to emit a file Word refuses to open."""
    from headnote.drafter import layout_presets as LP

    out = LP.sanitise(junk, "hi")
    if out is None:
        return
    pg = out["page"]
    assert 3.0 <= pg["width"] <= 24.0 and 3.0 <= pg["height"] <= 36.0
    for fmt in out["roles"].values():
        assert 4 <= fmt.get("size", 14) <= 72
    # and it must still render
    blocks = HB.from_document_html(TA.document("bail_sessions", {}, "hi"))
    assert _docx_facts(LT.render_into_layout(out, blocks))["paras"] > 10


def test_resolve_reports_the_format_the_file_actually_has(monkeypatch):
    """`X-Headnote-Format` is what the screen tells him he is holding. If he asks
    for his own format and it has gone, the honest answer is "standard" — not
    "own" over a file that is nothing of the kind."""
    from headnote.drafter import layout_presets as LP
    from headnote.drafter import style_profile as SP

    monkeypatch.setattr(SP, "load_layout", lambda uid: None)
    tpl, how = LP.resolve(layout_id="own", layout=None, user_id="u1", lang="hi")
    assert how == "standard" and tpl is not None

    monkeypatch.setattr(SP, "load_layout",
                        lambda uid: _capture("Headnote_Sample_bail_sessions.docx"))
    _, how = LP.resolve(layout_id="own", layout=None, user_id="u1", lang="hi")
    assert how == "own"
    _, how = LP.resolve(layout_id="legal_district", layout=None, user_id="u1", lang="hi")
    assert how == "legal_district"


# ---------------------------------------------------------------- the routes

def test_layout_presets_route_leads_with_his_own_format(client, monkeypatch):
    from headnote.drafter import style_profile as SP

    monkeypatch.setattr(SP, "load_layout", lambda uid: None)
    d = client.get("/api/draft/layout-presets?lang=hi").json()
    assert d["has_own"] is False
    assert [p["id"] for p in d["presets"]][0] == "standard"

    monkeypatch.setattr(SP, "load_layout",
                        lambda uid: _capture("Headnote_Sample_bail_sessions.docx"))
    d = client.get("/api/draft/layout-presets?lang=hi").json()
    assert d["has_own"] is True
    assert d["presets"][0]["id"] == "own", "his own filed paper must lead the rail"
    assert d["presets"][0]["font"] == "Mangal"
    # the browser draws a true-to-scale thumbnail, so it needs the real geometry
    assert d["presets"][0]["page"]["width"] == 8.5


def test_layout_presets_route_is_not_shadowed_by_the_draft_id_route(client, monkeypatch):
    """`GET /{draft_id}` is declared in the same router. Any new GET must sit
    ABOVE it or FastAPI treats the path as a draft id — which is exactly what
    happened to this one first time round."""
    from headnote.drafter import style_profile as SP

    monkeypatch.setattr(SP, "load_layout", lambda uid: None)
    r = client.get("/api/draft/layout-presets")
    assert r.status_code == 200, r.text
    assert "presets" in r.json(), "the route was swallowed by /{draft_id}"


@pytest.mark.parametrize("pid", ["standard", "legal_district", "hc_double", "compact"])
def test_export_honours_the_picked_layout(client, monkeypatch, pid):
    """The header names the layout, so the screen and the file can never
    disagree about what he is about to file."""
    from headnote.drafter import layout_presets as LP
    from headnote.drafter import style_profile as SP

    monkeypatch.setattr(SP, "load_layout",
                        lambda uid: _capture("Headnote_Sample_bail_sessions.docx"))
    r = client.post("/api/draft/template-docx",
                    json={"doc_type": "bail_sessions", "fields": {}, "lang": "hi",
                          "layout_id": pid})
    assert r.status_code == 200, r.text
    assert r.headers["X-Headnote-Format"] == pid
    assert _docx_facts(r.content)["page"] == _preset_facts(pid)["page"], \
        "he picked a layout and got a different page"


def test_export_honours_a_reference_layout(client, monkeypatch):
    from headnote.drafter import layout_presets as LP
    from headnote.drafter import style_profile as SP

    # he HAS his own DNA — the reference he chose for this matter must still win
    monkeypatch.setattr(SP, "load_layout",
                        lambda uid: _capture("Headnote_Sample_bail_sessions.docx"))
    with open("Headnote_Bail_Hindi.docx", "rb") as fh:
        ref = LP.capture_reference(fh.read())
    r = client.post("/api/draft/template-docx",
                    json={"doc_type": "bail_sessions", "fields": {}, "lang": "hi",
                          "layout_id": "own", "layout": ref})
    assert r.status_code == 200
    assert r.headers["X-Headnote-Format"] == "reference"
    assert _docx_facts(r.content)["fonts"] == {"Times New Roman"}


def test_reference_route_refuses_a_scan_instead_of_inventing_a_format(client):
    """A PDF or a photo has no layout stored inside it. Saying so plainly beats
    handing back a format we made up."""
    r = client.post("/api/draft/reference-format",
                    files={"file": ("order.pdf", b"%PDF-1.4 not really", "application/pdf")},
                    data={"lang": "hi"})
    assert r.status_code == 400
    assert ".docx" in r.json()["error"]


def test_reference_route_reads_a_real_docx(client):
    with open("Headnote_Bail_Hindi.docx", "rb") as fh:
        r = client.post("/api/draft/reference-format",
                        files={"file": ("ref.docx", fh.read(), DOCX_MIME)},
                        data={"lang": "hi"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["ok"] and d["font"] == "Times New Roman"
    assert d["page"]["width"] == 8.5
    assert d["layout"]["roles"], "the layout the browser sends back was not returned"
