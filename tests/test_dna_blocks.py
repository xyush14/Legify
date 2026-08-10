"""The DNA seam: authored payload → role-tagged blocks → the advocate's own .docx.

These guard the join that makes Draft DNA reachable from drafting at all. The
load-bearing one is `test_no_html_ever_reaches_the_docx` — the HTML preview marks
ungrounded facts and flagged citations with <span> wrappers, and a filed court
document must never contain them.
"""
import io

import pytest

from headnote.drafter import dna_blocks as DB
from headnote.drafter import layout_template as LT


def _payload():
    return {
        "side_label": "आवेदक की ओर से",
        "court_name": "न्यायालय श्रीमान अपर सत्र न्यायाधीश महोदय, भोपाल (म.प्र.)",
        "case_code": "जमानत आवेदन क्रमांक",
        "case_number": "512",
        "case_year": "2026",
        "applicant_label": "आवेदक",
        "applicant_desc": ["रमेश वर्मा पुत्र श्री मोहनलाल वर्मा, आयु 34 वर्ष",
                           "निवासी <u>14, शिवाजी नगर, भोपाल</u>"],
        "respondent_label": "अनावेदक",
        "respondent_desc": ["म.प्र. राज्य द्वारा थाना कोतवाली"],
        "versus": "विरुद्ध",
        "title_line": "धारा 483 बी.एन.एस.एस. के अंतर्गत जमानत आवेदन",
        "prelude": ["निवेदन इस प्रकार है:"],
        "paras": [
            {"kind": "ground", "text": "यह कि आवेदक दिनांक 12.03.2026 से अभिरक्षा में है।"},
            {"kind": "head", "text": "आधार"},
            {"kind": "ground", "text": "यह कि आवेदक का कोई पूर्व आपराधिक रिकॉर्ड नहीं है।"},
        ],
        "prayer": "अतः प्रार्थना है कि आवेदक को जमानत का लाभ दिया जावे।",
        "signatory_role": "आवेदक",
        "needs_verification": True,
        "verification": "मैं सत्यापित करता हूँ कि उपरोक्त कथन सत्य हैं।",
    }


def test_roles_are_in_document_order_and_known_to_the_layout_engine():
    blocks = DB.from_authored(_payload(), "hi")
    roles = [r for r, _ in blocks]
    # every role must be one the layout engine can format, else it silently
    # falls back to the default paragraph format
    assert set(roles) <= set(LT.ROLES), set(roles) - set(LT.ROLES)
    # header order mirrors _doc_header.render_header
    head = [r for r in roles if r in ("side_label", "court", "caseno", "applicant",
                                      "versus", "respondent", "section_heading")]
    assert head[:4] == ["side_label", "court", "caseno", "applicant"]
    assert head.index("versus") < head.index("respondent")
    # body then signature, never the other way round
    assert roles.index("ground") < roles.index("prayer") < roles.index("sig_party")


def test_grounds_are_numbered_and_a_head_does_not_advance_the_number():
    blocks = DB.from_authored(_payload(), "hi")
    grounds = [t for r, t in blocks if r == "ground"]
    assert grounds[0].startswith("1. ")
    assert grounds[1].startswith("2. ")          # the "आधार" head sat between them
    assert ("ground_head", "आधार") in blocks


def test_no_html_ever_reaches_the_docx():
    """Descriptor lines carry <u>…</u>, and the preview render injects grounding
    and citation markers. None of it may survive into a filed document."""
    p = _payload()
    p["paras"].append({"kind": "ground",
                       "text": 'भुगतान <span class="ung">₹5,00,000</span> का था &amp; शेष देय है'})
    blocks = DB.from_authored(p, "hi")
    joined = "\n".join(t for _, t in blocks)
    assert "<" not in joined and ">" not in joined
    assert "&amp;" not in joined and "&" in joined       # entity decoded, not dropped
    assert "14, शिवाजी नगर, भोपाल" in joined            # <u> stripped, text kept
    assert "₹5,00,000" in joined                        # the fact itself survives


def test_empty_payload_is_recognised_as_unrenderable():
    assert DB.is_empty(DB.from_authored({}, "hi"))
    assert not DB.is_empty(DB.from_authored(_payload(), "hi"))


# --------------------------------------------------------------------------
# the render half — blocks → a real .docx
# --------------------------------------------------------------------------
def _read(data: bytes):
    from docx import Document
    doc = Document(io.BytesIO(data))
    return doc, [p.text for p in doc.paragraphs]


def test_standard_format_when_the_advocate_has_no_dna(monkeypatch):
    """"Download .docx" must never be a dead action — no DNA still yields a
    filable document, just not in his own format."""
    from headnote.drafter import dna_layout, style_profile as SP
    pytest.importorskip("docx")
    monkeypatch.setattr(SP, "load_layout", lambda uid: None)

    data, how = dna_layout.render_blocks("user-1", DB.from_authored(_payload(), "hi"))
    assert how == "standard"
    doc, texts = _read(data)
    assert any("अपर सत्र न्यायाधीश" in t for t in texts)
    assert any(t.startswith("1. यह कि") for t in texts)
    # court name centred and bold in the standard format
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    court_p = next(p for p in doc.paragraphs if "अपर सत्र न्यायाधीश" in p.text)
    assert court_p.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert court_p.runs[0].bold is True


def test_the_advocates_own_geometry_and_font_are_reproduced(monkeypatch):
    from headnote.drafter import dna_layout, style_profile as SP
    pytest.importorskip("docx")
    tpl = {
        "page": {"width": 8.5, "height": 14.0, "margin_left": 1.2, "margin_right": 0.6,
                 "margin_top": 0.9, "margin_bottom": 0.7, "default_tab": 0.4},
        "roles": {
            "court":  {"align": "CENTER", "font": "Kruti Dev 010", "size": 18, "bold": True,
                       "line_spacing": 1.0, "tabs": []},
            "ground": {"align": "JUSTIFY", "font": "Kruti Dev 010", "size": 14,
                       "left_indent": 0.45, "line_spacing": 1.5, "tabs": []},
        },
    }
    monkeypatch.setattr(SP, "load_layout", lambda uid: tpl)

    data, how = dna_layout.render_blocks("user-1", DB.from_authored(_payload(), "hi"))
    assert how == "own"
    doc, _ = _read(data)
    sec = doc.sections[0]
    assert round(sec.page_width.inches, 2) == 8.5
    assert round(sec.page_height.inches, 2) == 14.0
    assert round(sec.left_margin.inches, 2) == 1.2

    ground = next(p for p in doc.paragraphs if p.paragraph_format.left_indent
                  and round(p.paragraph_format.left_indent.inches, 2) == 0.45)
    assert ground.runs, "the ground paragraph must carry text"
    # his own legacy font, and the text encoded into it (so it is NOT plain Unicode)
    assert ground.runs[0].font.name == "Kruti Dev 010"
    assert "यह कि" not in ground.text


# --------------------------------------------------------------------------
# UNIVERSALITY — this engine is for every Indian advocate, not one Hindi/Kruti Dev
# user. A Marathi or Tamil filing carrying a Hindi सत्यापन heading is a defective
# document, so these are correctness tests, not niceties.
# --------------------------------------------------------------------------
@pytest.mark.parametrize("lang,expect_place,expect_attest", [
    ("hi", "स्थान", "सत्यापन"),
    ("en", "Place", "VERIFICATION"),
    ("mr", "स्थळ", "सत्यापन"),          # from the reviewed court glossary
    ("gu", "સ્થળ", "સત્યાપન"),
])
def test_injected_labels_follow_the_advocates_language(lang, expect_place, expect_attest):
    blocks = DB.from_authored(_payload(), lang)
    joined = "\n".join(t for _, t in blocks)
    assert expect_place in joined
    assert expect_attest in joined


def test_an_unsupported_language_falls_back_to_english_never_to_hindi():
    """Tamil has no pinned glossary cells yet. English is neutral and every Indian
    advocate reads it; Hindi would be actively wrong on a Madras filing."""
    L = DB.labels("ta")
    assert L["place"] == "Place" and L["verification"] == "VERIFICATION"
    assert not any("ऀ" <= ch <= "ॿ" for ch in "".join(L.values())), \
        "no Devanagari may leak into a non-Hindi advocate's document"


def test_the_payload_can_override_every_injected_label():
    """The drafting brain authors in-language; when it names these blocks itself,
    its wording wins over any table of ours."""
    p = dict(_payload())
    p.update({"place_label": "இடம்", "date_label": "தேதி", "through_label": "வழக்கறிஞர் மூலம்",
              "advocate_label": "வழக்கறிஞர்", "verification_label": "உறுதிமொழி"})
    joined = "\n".join(t for _, t in DB.from_authored(p, "ta"))
    for word in ("இடம்", "தேதி", "வழக்கறிஞர் மூலம்", "உறுதிமொழி"):
        assert word in joined


@pytest.mark.parametrize("lang,expected_font", [
    ("en", "Times New Roman"),   # an English filing wants the court-standard serif
    ("hi", "Nirmala UI"),
    ("ta", "Nirmala UI"),        # one broad Unicode face covers every Indic script
])
def test_no_dna_fallback_font_suits_the_language(monkeypatch, lang, expected_font):
    from headnote.drafter import dna_layout, style_profile as SP
    pytest.importorskip("docx")
    monkeypatch.setattr(SP, "load_layout", lambda uid: None)
    data, how = dna_layout.render_blocks(None, DB.from_authored(_payload(), lang), lang=lang)
    assert how == "standard"
    doc, _ = _read(data)
    fonts = {r.font.name for p in doc.paragraphs for r in p.runs if r.font.name}
    assert fonts == {expected_font}, fonts


def test_english_grounds_are_detected_by_the_zero_cost_path():
    """The deterministic detector must not be Hindi-only — an English-drafting
    advocate's own filings have to capture without an LLM labeller too."""
    assert LT.detect_role("1. That the petitioner is innocent.", set()) == "ground"
    assert LT.detect_role("That, the accused was arrested on 12.03.2026.", set()) == "ground"


def test_dna_is_sticky_across_drafts_until_changed(monkeypatch):
    """Set once, applies to every draft thereafter — no re-upload, no per-draft step."""
    from headnote.drafter import dna_layout, style_profile as SP
    pytest.importorskip("docx")
    tpl = {"page": {"width": 8.5, "height": 14.0},
           "roles": {"ground": {"align": "JUSTIFY", "font": "Kruti Dev 010", "size": 14,
                                "line_spacing": 1.5, "tabs": []}}}
    calls = []

    def _load(uid):
        calls.append(uid)
        return tpl
    monkeypatch.setattr(SP, "load_layout", _load)

    for _ in range(3):
        _, how = dna_layout.render_blocks("uid-9", DB.from_authored(_payload(), "hi"))
        assert how == "own"
    assert calls == ["uid-9"] * 3      # re-read per draft, never cached-and-lost


def test_author_document_actually_hands_blocks_back(monkeypatch):
    """The point of the whole seam: the drafting brain's own return value carries the
    blocks, so nothing downstream has to re-derive them from HTML."""
    from headnote.drafter import author

    p = dict(_payload())
    p["_doc_type"] = "bail"
    monkeypatch.setattr(author, "author_payload", lambda *a, **k: p)

    out = author.author_document("रमेश वर्मा की जमानत", "bail", "hi")
    assert out["ok"] and out["html"]
    blocks = out["blocks"]
    assert not DB.is_empty(blocks)
    assert set(r for r, _ in blocks) <= set(LT.ROLES)
    # the HTML carries markup; the blocks must not
    assert "<" in out["html"] and "<" not in "".join(t for _, t in blocks)


def test_blocks_from_json_round_trip_are_accepted(monkeypatch):
    """Blocks are persisted as JSON, so they come back as lists, not tuples."""
    import json
    from headnote.drafter import dna_layout, style_profile as SP
    pytest.importorskip("docx")
    monkeypatch.setattr(SP, "load_layout", lambda uid: None)

    revived = json.loads(json.dumps(DB.from_authored(_payload(), "hi")))
    assert isinstance(revived[0], list)
    data, how = dna_layout.render_blocks(None, revived)
    assert how == "standard" and data[:2] == b"PK"
