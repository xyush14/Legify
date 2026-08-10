"""Draft DNA — the universal skeleton engine.

The promise: keep the advocate's own file as the template and change only what
varies. These tests assert the two things that make it universal — slots are
found by comparison alone (no language/layout knowledge), and filling never
disturbs the layout.
"""
import io

import pytest

from headnote.drafter import doc_skeleton as DS


def _doc(paras, font="Calibri"):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    for t in paras:
        r = d.add_paragraph().add_run(t)
        r.font.name = font
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_finds_slots_by_comparison_in_any_language():
    # Gujarati: identical boilerplate, one varying name — no rules, just diffing.
    a = _doc(["નામદાર કોર્ટ સમક્ષ", "અરજદાર — રમેશભાઈ પટેલ", "એ કે, અરજદાર નિર્દોષ છે."])
    b = _doc(["નામદાર કોર્ટ સમક્ષ", "અરજદાર — સુરેશભાઈ શાહ", "એ કે, અરજદાર નિર્દોષ છે."])
    skel = DS.build_skeleton([a, b])
    fixed = [x["fixed"] for x in skel["blocks"] if "fixed" in x]
    assert "નામદાર કોર્ટ સમક્ષ" in fixed          # recurs → his skeleton
    assert "એ કે, અરજદાર નિર્દોષ છે." in fixed
    assert DS.slot_count(skel) >= 1                # the differing name → a slot
    pat = [x for x in skel["blocks"] if "pattern" in x][0]
    assert any("અરજદાર" in p.get("lit", "") for p in pat["pattern"])   # his wording kept


def test_single_document_has_no_slots():
    skel = DS.build_skeleton([_doc(["एक", "दो"])])
    assert skel["n_docs"] == 1 and DS.slot_count(skel) == 0


def test_fill_replaces_only_the_slot_and_keeps_the_layout():
    docx = pytest.importorskip("docx")
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    def build(name):
        d = docx.Document()
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(3.0)
        p.paragraph_format.first_line_indent = Inches(-3.0)
        r = p.add_run(f"आवेदक ——— {name}")
        r.font.size = Pt(18)
        r.bold = True
        d.add_paragraph().add_run("यह कि आवेदक निर्दोष है।")
        b = io.BytesIO()
        d.save(b)
        return b.getvalue()

    spine, other = build("मोहर सिंह"), build("राजेन्द्र बघेल")
    skel = DS.build_skeleton([spine, other])
    keys = [k for _bi, _si, k, _p, _c in DS.iter_slots(skel)]
    assert keys, "expected at least one slot"
    out = DS.fill(spine, skel, {keys[0]: "राजू शर्मा"})

    a, b = docx.Document(io.BytesIO(spine)), docx.Document(io.BytesIO(out))
    assert len(a.paragraphs) == len(b.paragraphs)
    for pa, pb in zip(a.paragraphs, b.paragraphs):      # geometry must not drift
        assert pa.alignment == pb.alignment
        assert pa.paragraph_format.left_indent == pb.paragraph_format.left_indent
        assert pa.paragraph_format.first_line_indent == pb.paragraph_format.first_line_indent
    assert "राजू शर्मा" in "".join(r.text for r in b.paragraphs[0].runs)
    assert "यह कि आवेदक निर्दोष है।" in b.paragraphs[1].text   # fixed block untouched


def test_labels_and_field_mapping():
    a = _doc(["वादी — राम", "स्थायी पाठ"])
    b = _doc(["वादी — श्याम", "स्थायी पाठ"])
    skel = DS.label_slots(DS.build_skeleton([a, b]), use_llm=False)
    fields = DS.slot_fields(skel)
    assert fields and all("label" in f for f in fields)
    mapped = DS.fill_by_label(skel, {fields[0]["label"]: "मोहन"})
    assert mapped and list(mapped.values()) == ["मोहन"]
