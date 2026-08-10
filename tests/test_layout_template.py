"""Draft DNA — layout capture, merge, render, and id-keyed storage."""
import io

import pytest

from headnote.drafter import krutidev as kd
from headnote.drafter import layout_template as LT


def _docx(paras):
    """paras: list of (text, font_name). Build a minimal .docx in memory."""
    docx = pytest.importorskip("docx")
    doc = docx.Document()
    for text, font in paras:
        r = doc.add_paragraph().add_run(text)
        r.font.name = font
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_capture_tags_roles_and_geometry():
    # a tiny Kruti Dev cause-title + one ground
    data = _docx([
        (kd.to_krutidev("माननीय उच्च न्यायालय मध्य प्रदेश खण्डपीठ ग्वालियर"), "Kruti Dev 010"),
        (kd.to_krutidev("आवेदक ———\tराम कुमार पुत्र श्री श्याम"), "Kruti Dev 010"),
        (kd.to_krutidev("विरुद्ध"), "Kruti Dev 010"),
        (kd.to_krutidev("यहकि, आवेदक निर्दोष है।"), "Kruti Dev 010"),
    ])
    t = LT.capture_layout(data)
    assert "court" in t["roles"]
    assert "applicant" in t["roles"]
    assert "ground" in t["roles"]


def test_merge_prefers_recurring_geometry():
    a = {"page": {"width": 8.5}, "roles": {"ground": {"align": "JUSTIFY", "left_indent": None,
         "first_line_indent": None, "line_spacing": 1.5, "tabs": []}}}
    b = {"page": {"width": 8.5}, "roles": {"ground": {"align": "JUSTIFY", "left_indent": None,
         "first_line_indent": None, "line_spacing": 1.5, "tabs": []}}}
    c = {"page": {"width": 8.27}, "roles": {"ground": {"align": "LEFT", "left_indent": 1.0,
         "first_line_indent": None, "line_spacing": 1.0, "tabs": []}}}
    m = LT.merge_templates([a, b, c])
    assert m["page"]["width"] == 8.5          # modal page wins (2 of 3)
    assert m["roles"]["ground"]["align"] == "JUSTIFY"   # recurring format wins
    assert m["confidence"]["roles"]["ground"] == "2/3"


def test_render_reproduces_page_font_and_indents():
    tpl = {
        "page": {"width": 8.5, "height": 14.0, "margin_left": 1.5, "margin_right": 1.0,
                 "margin_top": 1.5, "margin_bottom": 1.0},
        "roles": {"applicant": {"align": "JUSTIFY", "left_indent": 3.0, "first_line_indent": -3.0,
                                "font": "Kruti Dev 010", "size": 18}},
    }
    data = LT.render_into_layout(tpl, [("applicant", "आवेदक ———\tराम कुमार")])
    docx = pytest.importorskip("docx")
    doc = docx.Document(io.BytesIO(data))
    sec = doc.sections[0]
    assert round(sec.page_width.inches, 1) == 8.5
    assert round(sec.page_height.inches, 1) == 14.0
    assert round(sec.left_margin.inches, 1) == 1.5
    p = doc.paragraphs[0]
    assert round(p.paragraph_format.left_indent.inches, 1) == 3.0
    assert round(p.paragraph_format.first_line_indent.inches, 1) == -3.0
    # text is emitted in his Kruti Dev font and decodes back to the Hindi
    assert p.runs[0].font.name == "Kruti Dev 010"
    assert kd.convert(p.runs[0].text) == "आवेदक ———\tराम कुमार"


def test_each_lawyer_renders_in_his_own_font():
    # A legacy-font (Kruti Dev / Devanagari) advocate and a Unicode-font
    # advocate must each get output in THEIR font — not one hard-coded font.
    kruti_tpl = {"page": {"width": 8.5}, "roles": {
        "ground": {"align": "JUSTIFY", "font": "Kruti Dev 010", "size": 16}}}
    guj_tpl = {"page": {"width": 8.27}, "roles": {
        "ground": {"align": "JUSTIFY", "font": "Shruti", "size": 14}}}

    dk = LT.render_into_layout(kruti_tpl, [("ground", "यह कि आवेदक निर्दोष है।")])
    dg = LT.render_into_layout(guj_tpl, [("ground", "એ કે અરજદાર નિર્દોષ છે.")])

    docx = pytest.importorskip("docx")
    pk = docx.Document(io.BytesIO(dk)).paragraphs[0].runs[0]
    pg = docx.Document(io.BytesIO(dg)).paragraphs[0].runs[0]
    # Kruti Dev advocate: font set to his, text ENCODED (decodes back to the Hindi)
    assert pk.font.name == "Kruti Dev 010"
    assert kd.convert(pk.text) == "यह कि आवेदक निर्दोष है।"
    # Unicode advocate: font set to his, Gujarati text passes through unchanged
    assert pg.font.name == "Shruti"
    assert pg.text == "એ કે અરજદાર નિર્દોષ છે."


def test_latin_and_punctuation_never_render_as_mojibake():
    """In a Kruti Dev run, 'IPC' would draw as 'प्च्ब्' and ':' as 'रु'. Such text
    must be peeled into a Latin-font run so the reader sees what we wrote."""
    tpl = {"page": {}, "roles": {"ground": {"align": "JUSTIFY", "font": "Kruti Dev 010", "size": 16}}}
    src = "यहकि, आवेदक के विरुद्ध IPC 420 (M.P.) प्रकरण क्रमांक: 124/2025 है।"
    data = LT.render_into_layout(tpl, [("ground", src)])
    docx = pytest.importorskip("docx")
    para = docx.Document(io.BytesIO(data)).paragraphs[0]
    # what a reader sees: Latin runs verbatim, Kruti runs decoded
    seen = "".join(r.text if r.font.name == "Times New Roman" else kd.convert(r.text)
                   for r in para.runs)
    assert seen == src
    assert any(r.font.name == "Times New Roman" for r in para.runs)   # Latin was peeled out


def test_unicode_font_is_never_split_or_encoded():
    segs = LT.split_script_segments("એ કે, અરજદાર vs રાજ્ય", "Shruti")
    assert len(segs) == 1 and segs[0][2] is False


def test_save_layout_aborts_when_the_read_fails(monkeypatch):
    """Read-modify-write on one jsonb column: a failed read must NOT overwrite,
    or a transient network error would destroy the advocate's text-style DNA."""
    import headnote.entitlements as ent
    from headnote.drafter import style_profile as SP

    class Flaky:
        def __init__(self):
            self.rows = {"u": {"format": {"para_prefix": "यहकि"}, "style_prose": "terse"}}

        def select(self, *a, **k):
            raise RuntimeError("network blip")

        def update(self, tbl, payload, params=None):
            self.rows[params["id"].split("eq.")[1]] = payload["draft_style"]

    sb = Flaky()
    monkeypatch.setattr(ent, "_supabase", sb)
    with pytest.raises(RuntimeError):
        SP.save_layout("u", {"page": {}, "roles": {"court": {}}})
    assert "format" in sb.rows["u"]          # untouched


def test_saving_text_style_preserves_the_layout(monkeypatch):
    import headnote.entitlements as ent
    from headnote.drafter import style_profile as SP

    class SB:
        def __init__(self):
            self.rows = {}

        def select(self, tbl, params=None):
            u = params["id"].split("eq.")[1]
            return [{"draft_style": self.rows[u]}] if u in self.rows else []

        def update(self, tbl, payload, params=None):
            self.rows[params["id"].split("eq.")[1]] = payload["draft_style"]

    monkeypatch.setattr(ent, "_supabase", SB())
    SP.save_layout("u", {"page": {"width": 8.5}, "roles": {"court": {"align": "CENTER"}}})
    SP.save_style("u", {"format": {"para_prefix": "यहकि"}})
    assert SP.load_layout("u") is not None                  # survived a text-style save
    SP.save_style("u", None)                                # clearing text DNA
    assert SP.load_layout("u") is not None                  # layout is separate DNA
    # and a layout-only profile must not read as "no DNA"
    assert SP.load_style("u") is not None


def test_layout_storage_is_keyed_to_user_id_and_isolated(monkeypatch):
    import headnote.entitlements as ent
    from headnote.drafter import style_profile as SP

    class FakeSB:
        def __init__(self):
            self.rows = {}

        def select(self, tbl, params=None):
            uid = params["id"].split("eq.")[1]
            return [{"draft_style": self.rows[uid]}] if uid in self.rows else []

        def update(self, tbl, payload, params=None):
            uid = params["id"].split("eq.")[1]
            self.rows[uid] = payload["draft_style"]

    monkeypatch.setattr(ent, "_supabase", FakeSB())

    tplA = {"page": {"width": 8.5}, "roles": {"court": {"align": "CENTER"}}}
    tplB = {"page": {"width": 8.27}, "roles": {"court": {"align": "LEFT"}}}
    # A already has text-style DNA; saving layout must MERGE, not clobber it
    SP.save_style("A", {"format": {"para_prefix": "यहकि"}})
    SP.save_layout("A", tplA)
    SP.save_layout("B", tplB)

    assert SP.load_layout("A")["page"]["width"] == 8.5
    assert SP.load_style("A") is not None                 # text DNA preserved
    assert SP.load_layout("B")["page"]["width"] == 8.27   # isolated per id
    assert SP.load_layout("A") != SP.load_layout("B")
    assert SP.load_layout("unknown-id") is None


# ---------------------------------------------------------------------------
# role detection against the shapes REAL filed drafts actually use. Each case
# below was a live mis-detection found by capturing an actual bail application:
# the wrong role means the wrong format lands on that block when we render.
# ---------------------------------------------------------------------------
def test_numbered_grounds_are_detected_as_grounds():
    """Filed grounds carry their own number. Before the prefix was stripped these
    fell through to the looser rules and became "dateline"/"court", which put body
    formatting on the cause-title."""
    for line in ("1. यहकि, पुलिस थाना कोतवाली द्वारा प्रार्थी को गिरफ्तार किया गया।",
                 "2. यह कि आवेदक निर्दोष है।",
                 "(3) यहकि, प्रकरण में विलम्ब है।",
                 "४। यह कि प्रार्थी स्थाई निवासी है।"):
        assert LT.detect_role(line, set()) == "ground", line


def test_a_ground_that_mentions_a_date_is_not_a_dateline():
    g = "5. यहकि, प्रार्थी दिनांक 12.03.2026 से न्यायिक अभिरक्षा में निरुद्ध है।"
    assert LT.detect_role(g, set()) == "ground"
    # the real dateline still is one
    assert LT.detect_role("दिनांक 12.03.2026", set()) == "dateline"
    assert LT.detect_role("स्थान — भोपाल", set()) == "dateline"


def test_a_ground_that_cites_a_court_is_not_the_court_line():
    long_ground = ("आवेदक का इस आशय का अन्य कोई जमानत आवेदन माननीय उच्चतम न्यायालय अथवा "
                   "माननीय उच्च न्यायालय में लंबित अथवा निराकृत नहीं है।")
    assert LT.detect_role(long_ground, set()) != "court"


def test_subordinate_court_lines_are_detected():
    """District filings address the forum as "न्यायालय श्रीमान …" — none of which
    contains "माननीय न्यायालय", so matching only that missed every one of them."""
    for line in ("न्यायालय माननीय सत्र न्यायाधीश महोदय, ग्वालियर (म.प्र.)",
                 "न्यायालय श्रीमान अपर सत्र न्यायाधीश महोदय, भोपाल (म.प्र.)",
                 "न्यायालय श्रीमान न्यायिक मजिस्ट्रेट प्रथम श्रेणी, विदिशा",
                 "माननीय उच्च न्यायालय मध्य प्रदेश खण्डपीठ ग्वालियर"):
        assert LT.detect_role(line, set()) == "court", line
    # still a once-only role
    assert LT.detect_role("न्यायालय श्रीमान सत्र न्यायाधीश, भोपाल", {"court"}) is None


def test_uncaptured_roles_fall_back_to_court_format_in_his_own_font():
    """His drafts won't contain every role. Those blocks must still look like a
    court document — and must stay in HIS typeface and size, not the default."""
    pytest.importorskip("docx")
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    tpl = {"page": {"width": 8.5, "height": 14.0},
           "roles": {"ground": {"align": "JUSTIFY", "font": "Kruti Dev 010", "size": 13,
                                "line_spacing": 1.5, "tabs": []}}}
    # ground_head / court were never captured
    data = LT.render_into_layout(tpl, [("court", "न्यायालय श्रीमान सत्र न्यायाधीश, भोपाल"),
                                       ("ground_head", "आधार"),
                                       ("ground", "1. यह कि आवेदक निर्दोष है।")])
    doc = Document(io.BytesIO(data))
    head = next(p for p in doc.paragraphs if p.runs and "vk/kkj" in p.text or "आधार" in p.text)
    assert head.alignment == WD_ALIGN_PARAGRAPH.CENTER   # not flat JUSTIFY
    assert head.runs[0].bold is True
    # his font and his dominant size carry onto the roles he never filed
    assert head.runs[0].font.name == "Kruti Dev 010"
    assert head.runs[0].font.size.pt == 13
