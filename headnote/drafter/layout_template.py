"""Draft DNA — LAYOUT TEMPLATE capture and reproduction.

The heart of Draft DNA. An advocate's filed .docx is his *layout fingerprint*:
page size, margins, and — per structural block (title, cause-title, grounds,
prayer, signature) — the exact alignment, indents, tab stops, spacing, font and
size. We capture that once, store it under his user id, and reproduce it exactly
when generating any new draft. Content is swapped; the layout is his.

Design for scale (lakhs of advocates):
  • NO per-user model. One shared, deterministic engine + a small per-user
    template (a few KB of JSON). Capture is one-time (on upload); generation is
    deterministic (no LLM for layout). Scales with ordinary web servers.
  • Per-user, id-keyed, isolated. The template is a plain dict, stored in
    user_profiles.draft_style["layout"] keyed by the advocate's user id
    (see style_profile.save_style / load_style).

Output text is encoded to the advocate's own font (Kruti Dev via krutidev.
to_krutidev) so the .docx opens looking like his own filing.
"""
from __future__ import annotations

import io
import logging
import re
from typing import Optional

from headnote.drafter import krutidev as _kd

log = logging.getLogger("headnote.drafter.layout")

# ---------------------------------------------------------------------------
# Structural ROLES a litigation draft is built from. Order = document order.
# Capture tags each of the advocate's paragraphs with one of these; generation
# emits these roles in a type-appropriate sequence, each in his captured format.
# ---------------------------------------------------------------------------
ROLES = (
    "side_label", "court", "caseno", "applicant", "versus", "respondent",
    "section_heading", "intro", "ground", "fact_head", "ground_head",
    "prayer", "dateline", "sig_party", "sig_by", "advocate",
)

_MAX_TABLES = 20   # bound what we persist per advocate (his index/record forms)

_APPLICANT_STARTS = ("आवेदक", "प्रार्थी", "वादी", "याचिकाकर्ता", "पुनरीक्षणकर्ता", "अपीलार्थी")
_RESPONDENT_STARTS = ("अनावेदक", "प्रत्यर्थी", "प्रतिवादी", "प्रतिपुनरीक्षणकर्ता", "अनावेदकगण")
# markers that a party line is a CAUSE-TITLE block (vs a bare signature designation)
_CAUSE_MARKERS = ("———", "——", "पुत्र", "निवासी", "पुलिस थाना", "\t")
_ONCE = ("court", "caseno", "section_heading", "intro", "fact_head",
         "ground_head", "prayer", "versus", "applicant", "respondent", "side_label")
# a filed ground carries its own number: "1.", "(2)", "३।", "4)" …
_NUM_PREFIX = re.compile(r"^[(\[]?\s*[0-9०-९]{1,3}\s*[.)\]:;।\-–]?\s*")
# How a numbered ground opens, per language. This is the DETERMINISTIC fallback
# only — `llm_labeller()` handles any language, including ones absent here; these
# anchors just make the zero-cost path work without a model. Add a language by
# adding its opener, nothing else.
_GROUND_STARTS = (
    "यहकि", "यह कि", "यह की",        # Hindi
    "हे की", "हेकी",                   # Marathi
    "કે,", "એ કે",                     # Gujarati
    "যে,",                             # Bengali
    "That ", "That,",                  # English
)


def _norm(txt: str) -> str:
    return _kd.convert(txt) if _kd.looks_like_krutidev(txt) else txt


def detect_role(decoded_text: str, seen: set) -> Optional[str]:
    """Best-effort role for one paragraph's DECODED text. Order matters — the
    same word (आवेदक) is a cause-title party at the top and a signature
    designation at the bottom, so specific/positional rules resolve first."""
    t = (decoded_text or "").strip()
    if not t:
        return None

    def once(role):  # header roles match only their FIRST occurrence
        return None if role in seen else role

    # 1) grounds — the repeating body. Filed drafts number them ("1. यहकि,",
    #    "(२) यह कि"), so the numbering prefix must be stripped before matching:
    #    without this the grounds — the bulk of the document — go undetected and
    #    then get mis-tagged by the looser rules below (a ground that mentions a
    #    date became a "dateline", one that cited माननीय उच्चतम न्यायालय became the
    #    "court"), which is how a body format ends up on the cause-title.
    if _NUM_PREFIX.sub("", t, count=1).startswith(_GROUND_STARTS):
        return "ground"
    # 2) signature side (specific, so it never eats the cause-title parties)
    if "बन्दी आवेदक" in t or "द्वारा अभिभाषक" in t or "द्वारा अधिवक्ता" in t:
        return "sig_by" if "अभिभाषक" in t or "अधिवक्ता" in t else "sig_party"
    if t in ("प्रार्थी", "याचिकाकर्ता", "पुनरीक्षणकर्ता", "आवेदक", "अपीलार्थी"):
        return "sig_party" if ("prayer" in seen or "ground_head" in seen) else None
    if ("एडवोकेट" in t or "अधिवक्ता" in t) and "द्वारा" not in t:
        return "advocate"
    # 3) singleton structural markers
    # a dateline IS the date line — not any paragraph that happens to mention a date
    if t.startswith(("दिनांक", "स्थान")) or (len(t) <= 40 and "दिनांक" in t):
        return "dateline"
    if t.startswith("अतः") or "निवेदन है कि" in t or "प्रार्थना है" in t:
        return once("prayer")
    if "संक्षिप्त तथ्य" in t or "तथ्य इस प्रकार" in t:
        return once("fact_head")
    if t.startswith("आधार"):
        return once("ground_head")
    if "निम्न प्रकार प्रस्तुत" in t or "निम्नानुसार" in t:
        return once("intro")
    if "अन्तर्गत धारा" in t or "प्रथम जमानत" in t or "याचिका अन्तर्गत" in t:
        return once("section_heading")
    if "विरुद्ध" in t or t.strip() == "बनाम":
        return once("versus")
    # the court line names the forum and is SHORT. Subordinate courts are addressed
    # "न्यायालय श्रीमान … महोदय" / "न्यायालय माननीय सत्र न्यायाधीश महोदय" — neither
    # contains "माननीय न्यायालय", so matching only that missed every district filing.
    # The length bound is what keeps a ground that cites a court from claiming it.
    if len(t) <= 90 and (t.startswith("न्यायालय")
                         or (t.startswith("माननीय") and ("न्यायाल" in t or "न्यायाधीश" in t))
                         # kept for court lines that open with something else
                         # ("समक्ष माननीय उच्च न्यायालय …"), now length-bounded so a
                         # ground reciting "…माननीय उच्च न्यायालय में लंबित नहीं है"
                         # can no longer claim the cause-title's format
                         or "माननीय उच्च न्यायालय" in t or "माननीय न्यायालय" in t):
        return once("court")
    if "की ओर से" in t:
        return once("side_label")
    if "एम-सी-आर-सी" in t or "क्रमांक" in t or "क्रमाकं" in t:
        return once("caseno")
    # 4) cause-title party blocks (only when the line looks like a party block)
    is_cause = any(m in t for m in _CAUSE_MARKERS)
    if is_cause and any(t.startswith(x) for x in _APPLICANT_STARTS):
        return once("applicant")
    if is_cause and any(t.startswith(x) for x in _RESPONDENT_STARTS):
        return once("respondent")
    return None


# ---------------------------------------------------------------------------
# CAPTURE — one advocate .docx → a role-tagged layout template
# ---------------------------------------------------------------------------
def _inches(v):
    try:
        return round(v.inches, 3)
    except Exception:
        return None


def _pt(v):
    try:
        return round(v.pt, 1)
    except Exception:
        return None


def _default_tab_inches(doc) -> float:
    """The document's `w:defaultTabStop` in inches (Word's own default is 0.5)."""
    try:
        from docx.oxml.ns import qn
        el = doc.settings.element.find(qn("w:defaultTabStop"))
        if el is not None:
            twips = int(el.get(qn("w:val")))
            return round(twips / 1440.0, 3)
    except Exception:
        pass
    return 0.5


def _para_format(p) -> dict:
    """Serialise a paragraph's exact geometry to plain JSON-safe values."""
    pf = p.paragraph_format
    fmt = {
        "align": (str(p.alignment).split()[0] if p.alignment is not None else None),
        "left_indent": _inches(pf.left_indent),
        "right_indent": _inches(pf.right_indent),
        "first_line_indent": _inches(pf.first_line_indent),
        "space_before": _pt(pf.space_before),
        "space_after": _pt(pf.space_after),
        "line_spacing": (round(pf.line_spacing, 2) if isinstance(pf.line_spacing, float) else None),
        "tabs": [],
        "font": None,
        "size": None,
        "bold": False,
    }
    try:
        for t in pf.tab_stops:
            fmt["tabs"].append({
                "pos": _inches(t.position),
                "align": str(t.alignment).split()[0] if t.alignment is not None else "LEFT",
            })
    except Exception:
        pass
    for r in p.runs:
        if r.text.strip():
            fmt["font"] = _kd._run_font(r) or fmt["font"]
            fmt["size"] = _pt(r.font.size) or fmt["size"]
            fmt["bold"] = bool(r.bold)
            break
    return fmt


def capture_layout(data: bytes | str, labeller=None) -> dict:
    """One advocate .docx → a layout template: page geometry + per-role format
    + tables. `labeller(decoded_paragraphs) -> [role|None]` decides each block's
    role; pass `llm_labeller(lang)` for any-language support, or leave None for
    the fast deterministic (Hindi-convention) labels. Font is reported per block
    (e.g. 'Kruti Dev 010' or 'Shruti') so generation honours his typeface.
    """
    from docx import Document

    doc = Document(io.BytesIO(data) if isinstance(data, (bytes, bytearray)) else data)
    sec = doc.sections[0]
    page = {
        "width": _inches(sec.page_width), "height": _inches(sec.page_height),
        "margin_left": _inches(sec.left_margin), "margin_right": _inches(sec.right_margin),
        "margin_top": _inches(sec.top_margin), "margin_bottom": _inches(sec.bottom_margin),
        # the default tab interval decides where an un-stopped tab lands, which is
        # how cause-titles push the party's name right — capture it, don't assume.
        "default_tab": _default_tab_inches(doc),
    }
    # decode every paragraph once (font-aware), then LABEL. Labelling is either
    # the LLM (general, any language) or the deterministic detector (fast, Hindi
    # district conventions) — the label source is pluggable via `labeller`.
    decoded = [_kd._decode_paragraph(p) for p in doc.paragraphs]
    labels = labeller(decoded) if labeller else _label_deterministic(decoded)

    # A filing often repeats its header blocks (an index sheet, then the main
    # application) at DIFFERENT sizes. We want the MAIN application's format, so
    # for each role pick the occurrence nearest the grounds rather than the first.
    paras = list(doc.paragraphs)
    ground_idxs = [i for i, r in enumerate(labels) if r == "ground"]
    anchor = ground_idxs[0] if ground_idxs else 0
    occurrences: dict[str, list[int]] = {}
    for i, role in enumerate(labels):
        if role:
            occurrences.setdefault(role, []).append(i)
    roles: dict[str, dict] = {}
    for role, idxs in occurrences.items():
        best = min(idxs, key=lambda i: abs(i - anchor))
        roles[role] = _para_format(paras[best])

    tables = []
    for t in doc.tables[:_MAX_TABLES]:
        widths = []
        try:
            for c in t.columns:
                widths.append(round(c.width.inches, 3) if c.width else None)
        except Exception:
            pass
        tables.append({"cols": len(t.columns), "rows": len(t.rows), "widths": widths})

    # The advocate's ACTUAL page, block by block, in document order. Used only to
    # show him a faithful reproduction on the confirm screen — it is NOT part of
    # the stored template (extract-then-discard: we persist geometry, never his
    # client's facts).
    preview = []
    for p, role in zip(paras, labels):
        t = _kd._decode_paragraph(p).rstrip()
        if not t.strip():
            continue
        preview.append({"role": role or "text", "text": t[:400],
                        "fmt": _para_format(p)})
        if len(preview) >= 60:
            break

    return {"page": page, "roles": roles, "roles_seen": sorted(roles.keys()),
            "tables": tables, "preview_blocks": preview}


def _label_deterministic(decoded: list[str]) -> list[Optional[str]]:
    """Fast label pass using the Hindi district-court anchors. Good default for
    the core market; the LLM labeller generalises to any language."""
    seen: set = set()
    out = []
    for dec in decoded:
        role = detect_role(dec, seen)
        if role:
            seen.add(role)
        out.append(role)
    return out


def llm_labeller(lang: str = "hi"):
    """A labeller that asks the LLM to tag each paragraph's structural ROLE — works
    in ANY language (Gujarati, Bengali, Marathi, Tamil, English). Returns a
    callable(decoded_paragraphs) -> [role|None]. Falls back to the deterministic
    labels for any paragraph the model leaves blank/unknown."""
    def _run(decoded: list[str]) -> list[Optional[str]]:
        fallback = _label_deterministic(decoded)
        numbered = "\n".join(f"{i}: {t.strip()[:200]}" for i, t in enumerate(decoded) if t.strip())
        if not numbered.strip():
            return fallback
        sys = (
            "You label the structural role of each paragraph of an Indian court "
            "application, in ANY language. Roles: " + ", ".join(ROLES) + ". Return ONLY "
            "JSON {\"labels\": {\"<index>\": \"<role>\"}} for paragraphs you can classify; "
            "omit ones you cannot. 'ground' = a numbered argument paragraph; 'applicant'/"
            "'respondent' = the cause-title party blocks; 'sig_party'/'advocate' = the "
            "signature block. Judge by MEANING and POSITION, not specific words."
        )
        try:
            from headnote.llm.client import _call_deepseek_or_groq, parse_json_response
            from headnote import config
            raw, _ = _call_deepseek_or_groq(
                sys, f"PARAGRAPHS:\n{numbered}\n\nReturn the JSON.",
                max_tokens=1500, claude_model=config.DRAFTER_AUTHOR_MODEL, json_mode=True,
            )
            parsed = parse_json_response(raw) or {}
            got = parsed.get("labels", parsed) if isinstance(parsed, dict) else {}
            out = list(fallback)
            # Only honour indices we actually OFFERED. Blank paragraphs are absent
            # from the prompt, so a model that renumbered sequentially would skew
            # every label onto the wrong paragraph — reject those instead.
            offered = {i for i, t in enumerate(decoded) if t.strip()}
            for k, v in (got.items() if isinstance(got, dict) else []):
                try:
                    idx = int(k)
                except (TypeError, ValueError):
                    continue
                if idx in offered and v in ROLES:
                    out[idx] = v
            return out
        except Exception:
            log.warning("llm_labeller failed — using deterministic labels", exc_info=True)
            return fallback
    return _run


# ---------------------------------------------------------------------------
# CONSENSUS — combine an advocate's 2–3 drafts into ONE template
# ---------------------------------------------------------------------------
def merge_templates(templates: list[dict]) -> dict:
    """Fold several per-draft templates into the advocate's consensus template.

    Page geometry: the modal value across drafts. Per-role format: the value
    that recurs across drafts wins (a habit), else the first seen. Records how
    many drafts each role/page value appeared in as confidence."""
    templates = [t for t in templates if isinstance(t, dict) and t.get("roles")]
    if not templates:
        return {"page": {}, "roles": {}, "confidence": {}}

    # page: modal per key
    page = {}
    pconf = {}
    for key in ("width", "height", "margin_left", "margin_right", "margin_top",
                "margin_bottom", "default_tab"):
        vals = [t["page"].get(key) for t in templates if t.get("page", {}).get(key) is not None]
        if vals:
            best = max(set(vals), key=vals.count)
            page[key] = best
            pconf[key] = f"{vals.count(best)}/{len(templates)}"

    # roles: modal format per role (compare by a hashable signature)
    roles = {}
    rconf = {}
    all_roles = {r for t in templates for r in t.get("roles", {})}
    for role in all_roles:
        variants = [t["roles"][role] for t in templates if role in t.get("roles", {})]
        # signature ignoring font/size drift we don't want to over-weight
        def sig(f):
            return (f.get("align"), f.get("left_indent"), f.get("first_line_indent"),
                    f.get("line_spacing"), tuple((x["pos"], x["align"]) for x in f.get("tabs", [])))
        sigs = [sig(v) for v in variants]
        best_sig = max(set(sigs), key=sigs.count)
        chosen = next(v for v in variants if sig(v) == best_sig)
        roles[role] = chosen
        rconf[role] = f"{sigs.count(best_sig)}/{len(templates)}"

    # tables: keep the richest table-set the advocate uses (the draft with the
    # most captured tables) — his index/record forms.
    tables = max((t.get("tables") or [] for t in templates), key=len, default=[])
    # the richest real page, for the confirm screen only (stripped before storage)
    preview = max((t.get("preview_blocks") or [] for t in templates), key=len, default=[])

    return {"page": page, "roles": roles, "tables": tables, "preview_blocks": preview,
            "confidence": {"page": pconf, "roles": rconf}, "n_drafts": len(templates)}


def strip_preview(tpl: dict) -> dict:
    """The persistable template — geometry only, never the advocate's own text."""
    return {k: v for k, v in (tpl or {}).items() if k != "preview_blocks"}


# ---------------------------------------------------------------------------
# Font registry — the ONLY per-language piece. A legacy ASCII-mapped font (like
# Kruti Dev for Devanagari) needs a Unicode→glyph encoder; a modern Unicode font
# (Nirmala/Shruti/Vrinda…) needs nothing (its text is already real). Register a
# language's legacy font here (Gujarati Shree/LMG, Bengali Bijoy…) and the whole
# engine supports it — nothing else changes.
# ---------------------------------------------------------------------------
# Each entry: (font detector, Unicode→glyph encoder, chars unsafe in that font).
LEGACY_ENCODERS = [
    (_kd.is_krutidev_font, _kd.to_krutidev, _kd.UNSAFE_IN_KRUTIDEV),   # Devanagari
]


def _legacy_spec(font_name: Optional[str]):
    """(encoder, unsafe_chars) for a legacy ASCII-mapped font, else (None, None)."""
    for detect, enc, unsafe in LEGACY_ENCODERS:
        if detect(font_name or ""):
            return enc, unsafe
    return None, None


def _legacy_encoder(font_name: Optional[str]):
    return _legacy_spec(font_name)[0]


def _encode_for_font(text: str, font_name: Optional[str]) -> str:
    """Encode text into the glyph bytes a legacy font expects; Unicode fonts
    pass through unchanged (already real text)."""
    enc = _legacy_encoder(font_name)
    return enc(text) if enc else text


# Text a legacy font cannot represent (Latin words like "IPC", and punctuation
# like ':' that the font would redraw as 'रु') must live in its own Latin-font
# run — exactly what advocates do by hand in their own filings. The unsafe set is
# derived from the font's mapping table, so it can't drift out of sync.
_LATIN_FALLBACK_FONT = "Times New Roman"


def split_script_segments(text: str, font_name: Optional[str]) -> list[tuple[str, str, bool]]:
    """[(segment_text, font_to_use, needs_legacy_encoding), …].

    Unicode font → one passthrough segment. Legacy font → maximal runs of
    encodable text (encoded, his font) interleaved with runs of unencodable text
    (verbatim, a Latin face) so nothing renders as mojibake.
    """
    enc, unsafe = _legacy_spec(font_name)
    if not enc or not text:
        return [(text, font_name or _LATIN_FALLBACK_FONT, False)]
    out: list[tuple[str, str, bool]] = []
    buf, buf_unsafe = [], False
    for ch in text:
        ch_unsafe = ch in unsafe
        if buf and ch_unsafe != buf_unsafe:
            out.append(("".join(buf), _LATIN_FALLBACK_FONT if buf_unsafe else font_name, not buf_unsafe))
            buf = []
        buf.append(ch)
        buf_unsafe = ch_unsafe
    if buf:
        out.append(("".join(buf), _LATIN_FALLBACK_FONT if buf_unsafe else font_name, not buf_unsafe))
    return out


def default_font_for(lang: Optional[str]) -> str:
    """The font to use for an advocate who has NOT captured his own yet.

    Deliberately just two answers rather than a per-language table we cannot
    justify: an English filing wants the court-standard serif, and every Indic
    script (Devanagari, Bengali, Gujarati, Gurmukhi, Tamil, Telugu, Kannada,
    Malayalam, Odia) is covered by one broad Unicode face. Once the advocate
    captures his DNA this is irrelevant — his own typeface is used instead.
    """
    return "Times New Roman" if (lang or "").strip().lower()[:2] == "en" else "Nirmala UI"


def standard_template(font: str = "Nirmala UI") -> dict:
    """A plain court-format layout for an advocate who has NOT captured his own yet.

    Draft DNA reproduces his measured geometry; this is the floor beneath it, so
    ".docx" is never an unavailable action — he gets a correct, filable document in
    standard shape, and the amber "your format isn't set up" line explains what he
    is missing. Deliberately conservative: A4, 1in margins, no invented flourishes.
    """
    def fmt(align, **kw):
        d = {"align": align, "left_indent": None, "right_indent": None,
             "first_line_indent": None, "space_before": None, "space_after": 6.0,
             "line_spacing": 1.5, "tabs": [], "font": font, "size": 14, "bold": False}
        d.update(kw)
        return d

    return {
        "page": {"width": 8.27, "height": 11.69, "margin_left": 1.0, "margin_right": 1.0,
                 "margin_top": 1.0, "margin_bottom": 1.0, "default_tab": 0.5},
        "roles": {
            "side_label":      fmt("RIGHT", size=13),
            "court":           fmt("CENTER", size=16, bold=True, space_after=10.0),
            "caseno":          fmt("CENTER", size=14, space_after=10.0),
            "applicant":       fmt("LEFT"),
            "versus":          fmt("CENTER", bold=True),
            "respondent":      fmt("LEFT", space_after=10.0),
            "section_heading": fmt("CENTER", bold=True, space_before=8.0, space_after=10.0),
            "intro":           fmt("JUSTIFY"),
            "fact_head":       fmt("CENTER", bold=True),
            "ground_head":     fmt("CENTER", bold=True),
            "ground":          fmt("JUSTIFY", first_line_indent=0.3),
            "prayer":          fmt("JUSTIFY", space_before=8.0),
            "dateline":        fmt("LEFT", space_after=2.0),
            "sig_party":       fmt("RIGHT", space_before=14.0, space_after=2.0),
            "sig_by":          fmt("RIGHT", space_after=2.0),
            "advocate":        fmt("RIGHT"),
            # ---- non-court families (doctypes.AUTHORITY_APPLICATION / NOTICE /
            # AFFIDAVIT / DEED). RENDER base formats only — deliberately NOT added to
            # ROLES, which is the CAPTURE vocabulary the labeller tags an advocate's
            # own .docx with. Reading his filings is therefore unchanged; these exist
            # so that when the parts renderer emits an addressee or a subject line,
            # `_role_base` has a sensible floor for it in HIS font and size instead of
            # the flat justified default.
            "addressee":       fmt("LEFT", space_after=2.0),
            "subject":         fmt("LEFT", bold=True, space_before=6.0, space_after=6.0),
            "reference":       fmt("LEFT", space_after=6.0),
            "salutation":      fmt("LEFT", space_before=6.0, space_after=8.0),
            "enclosures":      fmt("LEFT", space_after=2.0),
            "witness":         fmt("LEFT", space_after=6.0),
        },
    }


def template_primary_font(template: dict) -> str:
    """The advocate's dominant font across captured blocks (his typeface),
    defaulting to a broad Unicode Devanagari face when unknown."""
    fonts = [f.get("font") for f in (template or {}).get("roles", {}).values() if f.get("font")]
    return max(set(fonts), key=fonts.count) if fonts else "Nirmala UI"


def template_primary_size(template: dict) -> Optional[float]:
    """His dominant body size, so roles his drafts didn't contain don't come out
    at some unrelated default size next to the ones they did."""
    sizes = [f.get("size") for f in (template or {}).get("roles", {}).values()
             if isinstance(f.get("size"), (int, float))]
    return max(set(sizes), key=sizes.count) if sizes else None


def _role_base(role: str, primary_font: str, primary_size: Optional[float]) -> dict:
    """Format to start from for one role, BEFORE the advocate's captured format is
    laid over it. His own drafts rarely contain every role (a bail application has
    no `fact_head`), and falling back to one flat justified default put headings and
    signature lines in body format. So the floor is the standard court format for
    that role — restated in HIS typeface and size, which are the parts we do know."""
    base = standard_template(primary_font)["roles"].get(role)
    base = dict(base) if base else _default_fmt(primary_font)
    base["font"] = primary_font
    if primary_size:
        base["size"] = primary_size
    return base


# ---------------------------------------------------------------------------
# RENDER — fill the advocate's template with content → his-layout .docx
# ---------------------------------------------------------------------------
_ALIGN = None
_TAB = None


def _load_enums():
    global _ALIGN, _TAB
    if _ALIGN is None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        _ALIGN = {
            "LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
            "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT, "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        _TAB = {"LEFT": WD_TAB_ALIGNMENT.LEFT, "CENTER": WD_TAB_ALIGNMENT.CENTER,
                "RIGHT": WD_TAB_ALIGNMENT.RIGHT, "DECIMAL": WD_TAB_ALIGNMENT.DECIMAL}


def _default_fmt(primary_font: str) -> dict:
    return {"align": "JUSTIFY", "left_indent": None, "right_indent": None,
            "first_line_indent": None, "space_before": None, "space_after": None,
            "line_spacing": 1.5, "tabs": [], "font": primary_font, "size": 14, "bold": False}


def render_into_layout(template: dict, blocks: list[tuple]) -> bytes:
    """Build a .docx that reproduces the advocate's layout, IN HIS OWN FONT.

    `blocks` = list of (role, content). content is unicode text; if role ==
    'table', content is {"widths":[in…], "rows":[[cell,…],…]}. Each block is
    emitted with that role's captured formatting; text is encoded for the block's
    own font (legacy → glyph bytes, Unicode → passthrough) so it renders in the
    advocate's typeface, whatever language he drafts in.
    """
    from docx import Document
    from docx.shared import Inches, Pt

    _load_enums()
    tpl_roles = (template or {}).get("roles", {}) or {}
    page = (template or {}).get("page", {}) or {}
    primary = template_primary_font(template)
    primary_size = template_primary_size(template)

    out = Document()
    sec = out.sections[0]
    if page.get("width"):
        sec.page_width = Inches(page["width"])
    if page.get("height"):
        sec.page_height = Inches(page["height"])
    for attr, key in (("left_margin", "margin_left"), ("right_margin", "margin_right"),
                      ("top_margin", "margin_top"), ("bottom_margin", "margin_bottom")):
        if page.get(key) is not None:
            setattr(sec, attr, Inches(page[key]))
    # match his default tab interval — un-stopped tabs are what position the
    # party's name in a cause-title, so the wrong interval shifts the whole block
    if page.get("default_tab"):
        try:
            from docx.oxml.ns import qn
            settings = out.settings.element
            el = settings.find(qn("w:defaultTabStop"))
            if el is None:
                el = settings.makeelement(qn("w:defaultTabStop"), {})
                settings.append(el)
            el.set(qn("w:val"), str(int(round(page["default_tab"] * 1440))))
        except Exception:
            log.warning("could not set defaultTabStop", exc_info=True)

    for p in list(out.paragraphs):
        p._p.getparent().remove(p._p)

    for role, content in blocks:
        if role == "table":
            _emit_table(out, content, template, primary)
            continue
        fmt = _role_base(role, primary, primary_size)
        fmt.update(tpl_roles.get(role, {}) or {})   # his captured format always wins
        font = fmt.get("font") or primary
        p = out.add_paragraph()
        pf = p.paragraph_format
        if fmt.get("align") in _ALIGN:
            p.alignment = _ALIGN[fmt["align"]]
        if fmt.get("left_indent") is not None:
            pf.left_indent = Inches(fmt["left_indent"])
        if fmt.get("right_indent") is not None:
            pf.right_indent = Inches(fmt["right_indent"])
        if fmt.get("first_line_indent") is not None:
            pf.first_line_indent = Inches(fmt["first_line_indent"])
        if fmt.get("space_before") is not None:
            pf.space_before = Pt(fmt["space_before"])
        if fmt.get("space_after") is not None:
            pf.space_after = Pt(fmt["space_after"])
        if isinstance(fmt.get("line_spacing"), (int, float)):
            pf.line_spacing = fmt["line_spacing"]
        for t in fmt.get("tabs", []) or []:
            if t.get("pos") is not None:
                pf.tab_stops.add_tab_stop(Inches(t["pos"]), _TAB.get(t.get("align", "LEFT"), _TAB["LEFT"]))
        size = fmt.get("size") or 14
        bold = bool(fmt.get("bold"))
        for seg, seg_font, needs_enc in split_script_segments(str(content), font):
            if not seg:
                continue
            run = p.add_run(_encode_for_font(seg, font) if needs_enc else seg)
            _set_run_font(run, seg_font, size, bold)

    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def _emit_table(out, content: dict, template: dict, primary_font: str):
    """Recreate one of the advocate's tables — his column widths, new cell text."""
    from docx.shared import Inches, Pt

    widths = (content or {}).get("widths") or []
    rows = (content or {}).get("rows") or []
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    tbl = out.add_table(rows=0, cols=ncols)
    try:
        tbl.style = "Table Grid"
    except Exception:
        pass
    for r in rows:
        cells = tbl.add_row().cells
        for j in range(ncols):
            txt = str(r[j]) if j < len(r) else ""
            cell = cells[j]
            cell.text = ""
            para = cell.paragraphs[0]
            for seg, seg_font, needs_enc in split_script_segments(txt, primary_font):
                if not seg:
                    continue
                run = para.add_run(_encode_for_font(seg, primary_font) if needs_enc else seg)
                _set_run_font(run, seg_font, 12, False)
            if j < len(widths) and widths[j]:
                try:
                    cell.width = Inches(widths[j])
                except Exception:
                    pass


def _set_run_font(run, font_name: str, size_pt: float, bold: bool):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.bold = bold
    run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), font_name)
